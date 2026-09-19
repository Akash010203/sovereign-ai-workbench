"""rag/ingest.py — Document ingestion pipeline: text → chunks → embeddings → index."""
from __future__ import annotations

import logging
from pathlib import Path
from typing import Optional

from rag.chunking import chunk_text, TextChunk
from rag.embeddings import get_embedder
from rag.index import LocalVectorIndex

log = logging.getLogger(__name__)


class DocumentIngester:
    """
    Ingest documents into the local RAG knowledge base.

    Pipeline:
        file (txt/pdf) → extract text → chunk → embed → add to index
    """

    def __init__(self, index: LocalVectorIndex) -> None:
        self.index    = index
        self.embedder = get_embedder()

    def ingest_text(
        self,
        text: str,
        source: str,
        chunk_size: int = 400,
        overlap: int = 80,
    ) -> int:
        """
        Ingest a raw text string into the index.

        Returns the number of chunks added.
        """
        chunks = chunk_text(text, source, chunk_size, overlap)
        if not chunks:
            return 0
        texts      = [c.text for c in chunks]
        embeddings = self.embedder.embed_batch(texts)
        self.index.add_batch(chunks, embeddings)
        log.info("Ingested '%s': %d chunks", source, len(chunks))
        return len(chunks)

    def ingest_file(self, path: str | Path, chunk_size: int = 400) -> int:
        """
        Ingest a text document, office file, PDF, or OCR-readable image.

        PDFs are extracted with pdfminer when available; images use local OCR.
        """
        path = Path(path)
        source = path.name

        suffix = path.suffix.lower()
        if suffix == ".pdf":
            text = self._extract_pdf(path)
        elif suffix == ".docx":
            text = self._extract_docx(path)
        elif suffix == ".xlsx":
            text = self._extract_xlsx(path)
        elif suffix == ".pptx":
            text = self._extract_pptx(path)
        elif suffix in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}:
            text = self._extract_image(path)
        else:
            text = path.read_text(encoding="utf-8", errors="replace")

        return self.ingest_text(text, source=source, chunk_size=chunk_size)

    def ingest_directory(self, directory: str | Path, extensions: list[str] = None) -> int:
        """Ingest all matching files in a directory."""
        extensions = extensions or [".txt", ".pdf", ".md"]
        total = 0
        for file_path in Path(directory).rglob("*"):
            if file_path.suffix.lower() in extensions:
                try:
                    total += self.ingest_file(file_path)
                except Exception as exc:
                    log.warning("Failed to ingest %s: %s", file_path, exc)
        return total

    @staticmethod
    def _extract_pdf(path: Path) -> str:
        try:
            from pdfminer.high_level import extract_text
            return extract_text(str(path)) or ""
        except ImportError:
            log.warning("pdfminer not available — reading PDF as binary text fallback.")
            return path.read_bytes().decode("utf-8", errors="replace")

    @staticmethod
    def _extract_docx(path: Path) -> str:
        from docx import Document
        return "\n".join(paragraph.text for paragraph in Document(path).paragraphs)

    @staticmethod
    def _extract_xlsx(path: Path) -> str:
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        rows = []
        for sheet in workbook.worksheets:
            rows.append(f"Sheet: {sheet.title}")
            rows.extend(" | ".join(str(value) for value in row if value is not None)
                        for row in sheet.iter_rows(values_only=True))
        return "\n".join(row for row in rows if row)

    @staticmethod
    def _extract_pptx(path: Path) -> str:
        from pptx import Presentation
        return "\n".join(
            shape.text for slide in Presentation(path).slides
            for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()
        )

    @staticmethod
    def _extract_image(path: Path) -> str:
        """Read visible text from an image with the project's local OCR stack."""
        try:
            import pytesseract
            from PIL import Image
        except ImportError as exc:
            raise RuntimeError("Image OCR requires Pillow and pytesseract.") from exc

        text = pytesseract.image_to_string(Image.open(path), lang="eng").strip()
        if not text:
            raise RuntimeError("No readable text was found in this image.")
        return text
