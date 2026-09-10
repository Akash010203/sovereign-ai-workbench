"""tools/spreadsheet.py — Read/write Excel spreadsheets locally (openpyxl)."""
from __future__ import annotations

from pathlib import Path
from tools.registry import BaseTool, ToolResult


class SpreadsheetReadTool(BaseTool):
    @property
    def name(self) -> str: return "spreadsheet_read"
    @property
    def description(self) -> str:
        return "Read an Excel file and return its data as text. Args: path (str)."

    def run(self, path: str = "") -> ToolResult:
        file_path = Path(path)
        if not file_path.exists():
            return ToolResult(tool_name=self.name, success=False, output=None,
                              error=f"File not found: {path}")
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            output_lines = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                output_lines.append(f"=== Sheet: {sheet_name} ===")
                for row in ws.iter_rows(values_only=True):
                    output_lines.append("\t".join(str(c) if c is not None else "" for c in row))
            return ToolResult(tool_name=self.name, success=True,
                              output="\n".join(output_lines),
                              metadata={"path": path, "sheets": wb.sheetnames})
        except ImportError:
            return ToolResult(tool_name=self.name, success=False, output=None,
                              error="openpyxl not installed. Run: pip install openpyxl")
        except Exception as exc:
            return ToolResult(tool_name=self.name, success=False, output=None, error=str(exc))


class SpreadsheetWriteTool(BaseTool):
    @property
    def name(self) -> str: return "spreadsheet_write"
    @property
    def description(self) -> str:
        return ("Write tabular data to an Excel file. "
                "Args: path (str), rows (list[list]), headers (list[str]).")

    def run(self, path: str = "", rows: list = None, headers: list = None) -> ToolResult:
        try:
            import openpyxl
            from openpyxl.styles import Font
        except ImportError:
            return ToolResult(tool_name=self.name, success=False, output=None,
                              error="openpyxl not installed. Run: pip install openpyxl")

        out_path = Path(path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"

        if headers:
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True)

        for row in (rows or []):
            ws.append(list(row))

        wb.save(out_path)
        return ToolResult(tool_name=self.name, success=True,
                          output=f"Saved spreadsheet: {path}",
                          metadata={"path": path, "rows": len(rows or [])})
