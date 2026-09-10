"""tools/word.py — Generate Word (.docx) documents locally via python-docx."""
from __future__ import annotations

from pathlib import Path
from tools.registry import BaseTool, ToolResult


class WordWriteTool(BaseTool):
    """
    Create a Word document (.docx) from structured content.
    Uses python-docx — local only, no cloud.
    Install: pip install python-docx
    """

    @property
    def name(self) -> str:
        return "word_write"

    @property
    def description(self) -> str:
        return (
            "Generate a Word .docx document. "
            "Args: path (str), title (str), sections (list[dict with 'heading' and 'body'])."
        )

    def run(
        self,
        path: str = "output/document.docx",
        title: str = "Document",
        sections: list = None,
    ) -> ToolResult:
        try:
            from docx import Document
            from docx.shared import Pt
        except ImportError:
            return ToolResult(tool_name=self.name, success=False, output=None,
                              error="python-docx not installed. Run: pip install python-docx")

        doc = Document()

        # Title
        title_para = doc.add_heading(title, level=0)

        # Sections
        for section in (sections or []):
            heading = section.get("heading", "")
            body    = section.get("body", "")
            if heading:
                doc.add_heading(heading, level=1)
            if body:
                doc.add_paragraph(body)

        out_path = Path(path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(out_path)

        return ToolResult(
            tool_name=self.name, success=True,
            output=f"Word document saved: {path}",
            metadata={"path": path, "sections": len(sections or [])}
        )
