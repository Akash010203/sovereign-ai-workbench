"""tools/powerpoint.py — Generate PowerPoint (.pptx) presentations locally."""
from __future__ import annotations

from pathlib import Path
from tools.registry import BaseTool, ToolResult


class PowerPointWriteTool(BaseTool):
    """
    Create a PowerPoint presentation (.pptx) from slide data.
    Uses python-pptx — local only, no cloud.
    Install: pip install python-pptx
    """

    @property
    def name(self) -> str:
        return "powerpoint_write"

    @property
    def description(self) -> str:
        return (
            "Generate a PowerPoint .pptx file. "
            "Args: path (str), title (str), slides (list[dict with 'title' and 'content'])."
        )

    def run(
        self,
        path: str = "output/presentation.pptx",
        title: str = "Presentation",
        slides: list = None,
    ) -> ToolResult:
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
        except ImportError:
            return ToolResult(tool_name=self.name, success=False, output=None,
                              error="python-pptx not installed. Run: pip install python-pptx")

        prs = Presentation()
        slide_layout = prs.slide_layouts[1]  # title + content layout

        # Title slide
        title_slide_layout = prs.slide_layouts[0]
        title_slide = prs.slides.add_slide(title_slide_layout)
        title_slide.shapes.title.text = title
        if title_slide.placeholders[1:]:
            title_slide.placeholders[1].text = "SovereignAI — Generated Locally"

        # Content slides
        for slide_data in (slides or []):
            slide = prs.slides.add_slide(slide_layout)
            slide.shapes.title.text = slide_data.get("title", "")
            content = slide_data.get("content", "")
            if slide.placeholders[1:]:
                slide.placeholders[1].text = content

        out_path = Path(path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(out_path)

        return ToolResult(
            tool_name=self.name, success=True,
            output=f"PowerPoint saved: {path}",
            metadata={"path": path, "slides": len(slides or [])}
        )
